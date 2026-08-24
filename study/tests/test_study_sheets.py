"""Tests for the spreadsheet front end (export_study_sheet.py /
import_study_sheet.py) — an alternate way to collect the same
pre-registered study data run_study.py collects over the CLI. Covers the
things that matter most: no leakage of stimulus_id/kind/arm to the
participant (including across tabs in a batch workbook), that a filled-in
sheet round-trips into exactly the schema run_study.py itself writes, and
that a partially-completed batch only imports the tabs that are done.
"""

import json

import export_study_sheet
import import_study_sheet
from run_study import RESPONSE_LETTER_CODES
from stimuli import STIMULI

DATA_START_ROW = export_study_sheet.DATA_START_ROW


def _fill_ws(ws, responses, start_row=DATA_START_ROW):
    for i, (letter, confidence, warn) in enumerate(responses, start=start_row):
        ws.cell(row=i, column=4, value=letter)
        ws.cell(row=i, column=5, value=confidence)
        ws.cell(row=i, column=6, value=warn)


def _fill_sheet(path, responses, sheet_name="Responses"):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    _fill_ws(wb[sheet_name], responses)
    wb.save(path)


def _valid_responses():
    letters = list(RESPONSE_LETTER_CODES)  # C, R, I, V, S
    return [(letters[i % 5], (i % 5) + 1, "Y" if i % 2 else "N") for i in range(len(STIMULI))]


def test_export_does_not_leak_stimulus_id_kind_or_arm(tmp_path):
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("participant123", "treatment", cache)

    ws = wb["Responses"]
    visible_text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    instructions_text = "\n".join(
        str(c.value) for row in wb["Instructions"].iter_rows() for c in row if c.value is not None
    )

    # The literal id (e.g. "S1_digital_arrest") is the real leak risk — it
    # would tell the participant the message's category before they read
    # it. The word "scam" itself legitimately appears in genuine
    # explanation prose once a card is already shown, so that's not
    # checked here — only the structural id/kind/arm fields are.
    for stimulus in STIMULI:
        assert stimulus["id"] not in visible_text

    assert "control" not in visible_text.lower()
    assert "treatment" not in visible_text.lower()
    assert "control" not in instructions_text.lower()
    assert "treatment" not in instructions_text.lower()

    header = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert "arm" not in header
    assert "stimulus_id" not in header
    assert "kind" not in header


def test_export_row_order_matches_stimuli_order():
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("participant123", "control", cache)
    ws = wb["Responses"]
    rows = list(
        ws.iter_rows(
            min_row=DATA_START_ROW, max_row=DATA_START_ROW - 1 + len(STIMULI), values_only=True
        )
    )
    for i, (row, stimulus) in enumerate(zip(rows, STIMULI), start=1):
        assert row[0] == i
        assert row[1] == stimulus["text"]


def test_control_arm_never_shows_technique_cards():
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("p", "control", cache)
    ws = wb["Responses"]
    rows = ws.iter_rows(
        min_row=DATA_START_ROW, max_row=DATA_START_ROW - 1 + len(STIMULI), values_only=True
    )
    for row in rows:
        assert row[2] in ("This looks like a scam.", "No warning was flagged for this message.")


def test_round_trip_matches_run_study_schema(tmp_path, monkeypatch):
    pending = tmp_path / "pending"
    completed = tmp_path / "completed"
    responses_path = tmp_path / "responses.jsonl"
    monkeypatch.setattr(import_study_sheet, "PENDING_DIR", pending)
    monkeypatch.setattr(import_study_sheet, "COMPLETED_DIR", completed)
    monkeypatch.setattr(import_study_sheet, "RESPONSES_PATH", responses_path)

    pending.mkdir(parents=True)
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("abc123def456", "treatment", cache)
    sheet_path = pending / "study_abc123def456__treatment.xlsx"
    wb.save(sheet_path)
    _fill_sheet(sheet_path, _valid_responses())

    import_study_sheet.main()

    assert not sheet_path.exists()
    assert (completed / sheet_path.name).exists()

    lines = [json.loads(line) for line in responses_path.read_text(encoding="utf-8").splitlines()]
    assert len(lines) == len(STIMULI)
    for line, stimulus in zip(lines, STIMULI):
        assert set(line) == {
            "participant_id",
            "arm",
            "stimulus_id",
            "kind",
            "response",
            "confidence",
            "warn_others",
            "timestamp",
        }
        assert line["participant_id"] == "abc123def456"
        assert line["arm"] == "treatment"
        assert line["stimulus_id"] == stimulus["id"]
        assert line["kind"] == stimulus["kind"]
        assert line["response"] in RESPONSE_LETTER_CODES.values()
        assert 1 <= line["confidence"] <= 5
        assert isinstance(line["warn_others"], bool)


def test_incomplete_sheet_is_not_imported(tmp_path, monkeypatch):
    pending = tmp_path / "pending"
    completed = tmp_path / "completed"
    responses_path = tmp_path / "responses.jsonl"
    monkeypatch.setattr(import_study_sheet, "PENDING_DIR", pending)
    monkeypatch.setattr(import_study_sheet, "COMPLETED_DIR", completed)
    monkeypatch.setattr(import_study_sheet, "RESPONSES_PATH", responses_path)

    pending.mkdir(parents=True)
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("incomplete001", "control", cache)
    sheet_path = pending / "study_incomplete001__control.xlsx"
    wb.save(sheet_path)
    responses = _valid_responses()
    responses[-1] = ("", "", "")  # last row left blank
    _fill_sheet(sheet_path, responses)

    import_study_sheet.main()

    assert sheet_path.exists()  # left in place, not moved
    assert not responses_path.exists()


def test_duplicate_participant_is_skipped(tmp_path, monkeypatch):
    pending = tmp_path / "pending"
    completed = tmp_path / "completed"
    responses_path = tmp_path / "responses.jsonl"
    monkeypatch.setattr(import_study_sheet, "PENDING_DIR", pending)
    monkeypatch.setattr(import_study_sheet, "COMPLETED_DIR", completed)
    monkeypatch.setattr(import_study_sheet, "RESPONSES_PATH", responses_path)

    responses_path.write_text(
        json.dumps({"participant_id": "dupe001", "arm": "control", "stimulus_id": "x"}) + "\n",
        encoding="utf-8",
    )

    pending.mkdir(parents=True)
    cache = export_study_sheet.load_cache()
    wb = export_study_sheet.build_sheet("dupe001", "control", cache)
    sheet_path = pending / "study_dupe001__control.xlsx"
    wb.save(sheet_path)
    _fill_sheet(sheet_path, _valid_responses())

    before = responses_path.read_text(encoding="utf-8")
    import_study_sheet.main()
    after = responses_path.read_text(encoding="utf-8")

    assert sheet_path.exists()  # skipped, left in place
    assert before == after


def test_batch_workbook_does_not_leak_arm_or_id_in_tab_names_and_hides_index():
    cache = export_study_sheet.load_cache()
    assignments = [("pid_aaa111", "control"), ("pid_bbb222", "treatment"), ("pid_ccc333", "control")]
    wb = export_study_sheet.build_batch_workbook(assignments, cache)

    visible_sheetnames = [s for s in wb.sheetnames if wb[s].sheet_state != "hidden"]
    assert visible_sheetnames == ["Instructions", "P01", "P02", "P03"]
    for name in visible_sheetnames:
        assert "control" not in name.lower()
        assert "treatment" not in name.lower()
        for pid, _ in assignments:
            assert pid not in name

    assert "_index" in wb.sheetnames
    assert wb["_index"].sheet_state == "hidden"


def test_batch_round_trip_imports_all_tabs(tmp_path, monkeypatch):
    pending = tmp_path / "pending"
    completed = tmp_path / "completed"
    responses_path = tmp_path / "responses.jsonl"
    monkeypatch.setattr(import_study_sheet, "PENDING_DIR", pending)
    monkeypatch.setattr(import_study_sheet, "COMPLETED_DIR", completed)
    monkeypatch.setattr(import_study_sheet, "RESPONSES_PATH", responses_path)

    pending.mkdir(parents=True)
    cache = export_study_sheet.load_cache()
    assignments = [("batch_p1", "control"), ("batch_p2", "treatment")]
    wb = export_study_sheet.build_batch_workbook(assignments, cache)
    sheet_path = pending / "study_batch_test1234.xlsx"
    wb.save(sheet_path)

    wb2 = __import__("openpyxl").load_workbook(sheet_path)
    _fill_ws(wb2["P01"], _valid_responses())
    _fill_ws(wb2["P02"], _valid_responses())
    wb2.save(sheet_path)

    import_study_sheet.main()

    assert not sheet_path.exists()
    assert (completed / sheet_path.name).exists()

    lines = [json.loads(line) for line in responses_path.read_text(encoding="utf-8").splitlines()]
    assert len(lines) == 2 * len(STIMULI)
    p1_lines = [line for line in lines if line["participant_id"] == "batch_p1"]
    p2_lines = [line for line in lines if line["participant_id"] == "batch_p2"]
    assert len(p1_lines) == len(STIMULI)
    assert len(p2_lines) == len(STIMULI)
    assert all(line["arm"] == "control" for line in p1_lines)
    assert all(line["arm"] == "treatment" for line in p2_lines)


def test_batch_partial_completion_imports_only_finished_tabs(tmp_path, monkeypatch):
    pending = tmp_path / "pending"
    completed = tmp_path / "completed"
    responses_path = tmp_path / "responses.jsonl"
    monkeypatch.setattr(import_study_sheet, "PENDING_DIR", pending)
    monkeypatch.setattr(import_study_sheet, "COMPLETED_DIR", completed)
    monkeypatch.setattr(import_study_sheet, "RESPONSES_PATH", responses_path)

    pending.mkdir(parents=True)
    cache = export_study_sheet.load_cache()
    assignments = [("done_p1", "control"), ("unfinished_p2", "treatment")]
    wb = export_study_sheet.build_batch_workbook(assignments, cache)
    sheet_path = pending / "study_batch_partial01.xlsx"
    wb.save(sheet_path)

    from openpyxl import load_workbook

    wb2 = load_workbook(sheet_path)
    _fill_ws(wb2["P01"], _valid_responses())  # P02 left blank
    wb2.save(sheet_path)

    import_study_sheet.main()

    assert sheet_path.exists()  # left in pending — one tab still incomplete
    lines = [json.loads(line) for line in responses_path.read_text(encoding="utf-8").splitlines()]
    assert {line["participant_id"] for line in lines} == {"done_p1"}

    # re-running doesn't duplicate the already-imported participant
    import_study_sheet.main()
    lines_again = [json.loads(line) for line in responses_path.read_text(encoding="utf-8").splitlines()]
    assert len(lines_again) == len(lines)
