"""Export the gold-labeling sample to an .xlsx sheet for manual annotation.

Same sample as `label.py` (target=250, seed=42), so this is an alternate,
spreadsheet-based front end for the exact same gold set — not a different
sample. Existing answers in `labels_{annotator}.jsonl` are pre-filled with an
"x" in the matching technique column so you don't re-answer from scratch;
unanswered rows are left blank. A `notes` column carries a short list of
specific rows flagged during a manual review of the first-pass answers —
each is a suggestion to reconsider, not an automatic correction.

Default usage exports the full 250-message primary set, exactly as before.
Pass --annotator and --kappa-subset together to produce a small,
easy-to-hand-off sheet for a second annotator (see the plan's Phase A4 /
eval/kappa.py) — e.g. a friend labeling 80 messages in a spreadsheet is a
much smaller ask than cloning the repo and running a CLI script:

    python eval/export_labeling_sheet.py --annotator second --kappa-subset 80

Run: python eval/export_labeling_sheet.py
Then fill in backend/data/corpus/gold/gold_labeling.xlsx by hand, and run
eval/import_labeling_sheet.py to turn it back into labels_primary.jsonl.
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.taxonomy import TECHNIQUE_DESCRIPTIONS, Technique
from eval.label import LABELS_DIR, build_sample, load_corpus, load_existing_labels

TARGET = 250
SEED = 42

TECHNIQUES = list(Technique)

# Specific rows flagged during manual review of a first-pass answer.
# The batch reviewed 2026-08-04 was resolved (13 auto-applied, 5 re-answered
# by hand in gold_labelling_remaining.xlsx) — this is empty until the next
# review pass flags something new.
REVIEW_NOTES: dict[str, str] = {}

HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(bold=True)


def build_labels_sheet(wb: Workbook, sample: list[dict], existing: dict[str, list[str]]) -> None:
    ws = wb.active
    ws.title = "Labels"

    headers = ["id", "source", "text", "none"] + [t.value for t in TECHNIQUES] + ["notes"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for record in sample:
        reviewed = record["id"] in existing
        labels = existing.get(record["id"]) or []
        row = [record["id"], record["source"], record["text"]]
        row.append("x" if reviewed and not labels else "")
        row += ["x" if t.value in labels else "" for t in TECHNIQUES]
        row.append(REVIEW_NOTES.get(record["id"], ""))
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[3:-1]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 8
    for i in range(len(TECHNIQUES)):
        ws.column_dimensions[get_column_letter(5 + i)].width = 12
    ws.column_dimensions[get_column_letter(5 + len(TECHNIQUES))].width = 45

    ws.freeze_panes = "E2"
    ws.row_dimensions[1].height = 45


def build_legend_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Legend")
    ws.append(["technique", "description"])
    for col in (1, 2):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for t in TECHNIQUES:
        ws.append([t.value, TECHNIQUE_DESCRIPTIONS[t]])

    ws.append([])
    ws.append(["How to use the Labels sheet", ""])
    ws.append(["Put an 'x' in every technique column that applies to that message.", ""])
    ws.append(["If NO technique applies, mark the 'none' column with an 'x' instead.", ""])
    ws.append(["A row left completely blank means 'not reviewed yet', not 'nothing found' — it will be skipped on import.", ""])
    ws.append(["The 'notes' column carries suggestions from a prior review — read, then decide; not automatic.", ""])
    ws.append(["Common mix-ups:", ""])
    ws.append(["manufactured_urgency (time deadline) vs fake_scarcity (quantity/slots limited)", ""])
    ws.append(["false_authority (impersonates an institution) vs trust_transfer (impersonates a specific known person)", ""])
    ws.append(["reciprocity_hook: a bare prize-claim ('you've won, call to claim a £2000 bonus') counts on its own — you do NOT need an explicit 'do me a favor in return' ask for this to apply.", ""])

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--annotator", default="primary", help="Label suffix, e.g. 'primary' or 'second'")
    parser.add_argument(
        "--kappa-subset",
        type=int,
        default=None,
        help="Only export the first N messages of the sample (matches eval/label.py's --kappa-subset)",
    )
    args = parser.parse_args()

    is_default = args.annotator == "primary" and args.kappa_subset is None
    output_path = LABELS_DIR / ("gold_labeling.xlsx" if is_default else f"gold_labeling_{args.annotator}.xlsx")
    existing_path = LABELS_DIR / f"labels_{args.annotator}.jsonl"

    records = load_corpus()
    sample = build_sample(records, TARGET, SEED)
    if args.kappa_subset:
        sample = sample[: args.kappa_subset]
    existing = load_existing_labels(existing_path)

    wb = Workbook()
    build_labels_sheet(wb, sample, existing)
    build_legend_sheet(wb)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    prefilled = sum(1 for r in sample if r["id"] in existing)
    flagged = sum(1 for r in sample if r["id"] in REVIEW_NOTES)
    print(f"Wrote {output_path}")
    print(f"Sample size: {len(sample)}  |  Pre-filled from existing answers: {prefilled}  |  Flagged for review: {flagged}")


if __name__ == "__main__":
    main()
